import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def client(mini_env, monkeypatch):
    monkeypatch.setenv("AAPURTI_DB", str(mini_env["db"]))
    from aapurti.api.main import app
    return TestClient(app)


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["synthetic_data"] is True


def test_filters(client):
    j = client.get("/api/filters").json()
    assert len(j["plants"]) == 3 and j["skus"] and "SYNTHETIC" in j["notice"]


def test_kpis(client):
    j = client.get("/api/kpis").json()
    assert 0 < j["forecast_accuracy"] < 1 and abs(j["forecast_accuracy"] + j["forecast_wape"] - 1) < 1e-9
    assert j["tracked_sku_plants"] > 0 and 0 <= j["utilization_latest_month"] <= 1
    assert 0 <= j["dq_quarantine_rate"] < 0.1 and j["synthetic"] is True


def test_forecast_aggregate_and_filtered(client):
    j = client.get("/api/forecast").json()
    assert j["series"] and any(p["forecast"] is not None for p in j["series"])
    f = client.get("/api/filters").json()["skus"][0]
    j2 = client.get("/api/forecast", params={"sku_id": f["sku_id"], "plant_id": f["plant_id"]}).json()
    assert j2["series"] and j2["wape_holdout"] is not None
    assert sum(p["actual"] or 0 for p in j2["series"]) < sum(p["actual"] or 0 for p in j["series"])


def test_forecast_unknown_sku_is_empty_not_error(client):
    j = client.get("/api/forecast", params={"sku_id": "NOPE"}).json()
    assert j["series"] == []


def test_inventory_risk_sorted_and_filtered(client):
    rows = client.get("/api/inventory/risk").json()["rows"]
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    assert [order[r["stockout_risk"]] for r in rows] == sorted(order[r["stockout_risk"]] for r in rows)
    assert all(r["stockout_risk"] == "HIGH" for r in client.get("/api/inventory/risk", params={"risk": "high"}).json()["rows"])
    assert client.get("/api/inventory/risk", params={"risk": "bogus"}).status_code == 422
    assert len(client.get("/api/inventory/risk", params={"limit": 2}).json()["rows"]) <= 2


def test_capacity(client):
    j = client.get("/api/capacity", params={"plant_id": "P01"}).json()
    assert j["monthly"] and all(r["plant_id"] == "P01" for r in j["monthly"])
    assert all(0 <= r["utilization"] <= 1 for r in j["latest_by_line"])


def test_data_quality(client):
    j = client.get("/api/data-quality").json()
    assert j["reconciled"] is True and j["quarantine"] and len(j["row_counts"]) == 3
    for rc in j["row_counts"]:
        assert rc["bronze_rows"] == rc["silver_rows"] + rc["quarantined_rows"]


def test_backtest(client):
    j = client.get("/api/backtest").json()
    assert {r["model"] for r in j["results"]} >= {"naive", "gradient_boosting"} and j["winners"]


def test_export_xlsx(client):
    r = client.get("/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Summary" and {"Plant P01", "Plant P02", "Plant P03"} <= set(wb.sheetnames)


def test_missing_db_gives_503(monkeypatch, tmp_path):
    monkeypatch.setenv("AAPURTI_DB", str(tmp_path / "nope.db"))
    from aapurti.api.main import app
    assert TestClient(app).get("/api/kpis").status_code == 503
