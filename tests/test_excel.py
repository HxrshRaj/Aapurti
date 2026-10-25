"""Excel workbook: real formulas, and values after a LibreOffice headless recalculation match the pipeline."""
import io
import shutil
import subprocess

import pandas as pd
import pytest
from openpyxl import load_workbook

from aapurti.excel import build_workbook


@pytest.fixture(scope="module")
def tables(mini_env):
    import sqlite3
    con = sqlite3.connect(mini_env["db"])
    inv = pd.read_sql_query("SELECT * FROM fact_inventory_position", con)
    fva = pd.read_sql_query("SELECT * FROM fact_forecast", con)
    plants = pd.read_sql_query("SELECT plant_id, plant_name FROM dim_plant ORDER BY plant_id", con)
    con.close()
    return inv, fva, plants


@pytest.fixture(scope="module")
def xlsx_bytes(tables):
    inv, fva, plants = tables
    return build_workbook(inv, fva, plants, "2025-03-31")


def test_workbook_structure_and_formulas(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Summary", "Plant P01", "Plant P02", "Plant P03"]
    ws = wb["Plant P01"]
    assert ws["H5"].value.startswith("=IF(") and ws["J5"].value == "=F5*G5+I5"
    assert ws["K5"].value.startswith("=IF(") and ws["N5"].value == "=L5-M5"
    summary = wb["Summary"]
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 30 and any("COUNTIF(" in f for f in formulas) and any("SUMIF(" in f for f in formulas)
    assert any(f.startswith("=SUM(") for f in formulas)
    assert len(ws.conditional_formatting) >= 3


def test_no_hardcoded_totals(xlsx_bytes):
    ws = load_workbook(io.BytesIO(xlsx_bytes))["Plant P01"]
    total_row = next(r for r in range(5, 200) if ws.cell(r, 1).value == "Total")
    for col in (5, 12, 13, 14):
        assert str(ws.cell(total_row, col).value).startswith("=SUM(")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_libreoffice_recalculation_matches_pipeline(tmp_path, xlsx_bytes, tables):
    inv, fva, plants = tables
    src = tmp_path / "in.xlsx"
    src.write_bytes(xlsx_bytes)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180,
                   env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"})
    wb = load_workbook(out / "in.xlsx", data_only=True)

    for p in plants.itertuples():
        ws = wb[f"Plant {p.plant_id}"]
        mine = inv[inv["plant_id"] == p.plant_id].sort_values("sku_id").reset_index(drop=True)
        for i, row in mine.iterrows():
            r = 5 + i
            assert ws.cell(r, 1).value == row["sku_id"]
            doc = ws.cell(r, 8).value
            if pd.isna(row["days_of_cover"]):
                assert doc in (None, "")
            else:
                assert doc == pytest.approx(row["days_of_cover"], rel=1e-9)
            assert ws.cell(r, 10).value == pytest.approx(row["reorder_point"], rel=1e-9)
            assert ws.cell(r, 11).value == row["stockout_risk"], (p.plant_id, row["sku_id"])

    sm = wb["Summary"]
    assert sm["C5"].value == pytest.approx(inv[inv["plant_id"] == plants.plant_id[0]]["on_hand"].sum())
    total_row = next(r for r in range(5, 20) if sm.cell(r, 1).value == "Total")
    assert sm.cell(total_row, 3).value == pytest.approx(inv["on_hand"].sum())
    assert sm.cell(total_row, 5).value == (inv["stockout_risk"] == "HIGH").sum()
    assert sm.cell(total_row, 4).value == pytest.approx(inv["days_of_cover"].mean(), rel=1e-9)
    hold = fva[~fva["is_future"].astype(bool)]
    assert sm.cell(total_row, 6).value == pytest.approx(hold["actual"].sum())
    assert sm.cell(total_row, 7).value == pytest.approx(hold["forecast"].sum(), abs=len(inv) * 0.01)
    errors = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
    assert not errors
