"""Excel workbook (openpyxl) with live formulas: per-plant sheets, summary, pivot-style matrices.

Inputs/hard values (on-hand, demand, lead time, safety stock, forecast, actual) come from the pipeline;
days of cover, reorder point, risk, variance and all totals are Excel formulas so the workbook recalculates.
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import SYNTHETIC_NOTICE

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FIRST_ROW = 5
COLS = ["SKU", "Name", "Category", "Class", "On hand", "Avg daily demand (28d)", "Lead time (d)",
        "Days of cover", "Safety stock", "Reorder point", "Stockout risk",
        "Actual units (last 28d)", "Forecast units (last 28d)", "Variance (A-F)", "Variance %"]
# column letters used in formulas
C = {n: get_column_letter(i + 1) for i, n in enumerate(
    ["sku", "name", "cat", "cls", "oh", "avg", "lt", "doc", "ss", "rop", "risk", "act", "fc", "var", "varp"])}


def _header(ws, row, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _plant_sheet(wb, plant_id, plant_name, inv: pd.DataFrame, hold: pd.DataFrame) -> tuple[str, int, int]:
    title = f"Plant {plant_id}"
    ws = wb.create_sheet(title)
    ws["A1"] = f"{plant_id} {plant_name}: inventory position and last-28-day forecast vs actual"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = SYNTHETIC_NOTICE
    ws["A2"].font = Font(italic=True, color="7F7F7F")
    _header(ws, FIRST_ROW - 1, COLS)
    rows = inv[inv["plant_id"] == plant_id].sort_values("sku_id")
    h = hold.set_index("sku_id") if len(hold) else pd.DataFrame(columns=["actual", "forecast"])
    r = FIRST_ROW
    for x in rows.itertuples():
        ws.cell(r, 1, x.sku_id); ws.cell(r, 2, x.sku_name); ws.cell(r, 3, x.category); ws.cell(r, 4, x.volume_class)
        ws.cell(r, 5, float(x.on_hand)); ws.cell(r, 6, float(x.avg_daily_demand)); ws.cell(r, 7, int(x.lead_time_days))
        ws.cell(r, 8, f'=IF({C["avg"]}{r}>0,{C["oh"]}{r}/{C["avg"]}{r},"")')
        ws.cell(r, 9, float(x.safety_stock))
        ws.cell(r, 10, f'={C["avg"]}{r}*{C["lt"]}{r}+{C["ss"]}{r}')
        ws.cell(r, 11, f'=IF({C["doc"]}{r}="","LOW",IF({C["doc"]}{r}<{C["lt"]}{r},"HIGH",'
                       f'IF({C["oh"]}{r}<{C["rop"]}{r},"MEDIUM","LOW")))')
        a = float(h.loc[x.sku_id, "actual"]) if x.sku_id in h.index else 0.0
        f = float(h.loc[x.sku_id, "forecast"]) if x.sku_id in h.index else 0.0
        ws.cell(r, 12, a); ws.cell(r, 13, round(f, 2))
        ws.cell(r, 14, f'={C["act"]}{r}-{C["fc"]}{r}')
        ws.cell(r, 15, f'=IF({C["act"]}{r}=0,0,{C["var"]}{r}/{C["act"]}{r})')
        for col, fmt in ((5, "#,##0"), (6, "#,##0.0"), (8, "0.0"), (9, "#,##0.0"), (10, "#,##0.0"),
                         (12, "#,##0"), (13, "#,##0"), (14, "#,##0"), (15, "0.0%")):
            ws.cell(r, col).number_format = fmt
        r += 1
    if r == FIRST_ROW:      # plant with no SKUs: keep one blank row so ranges and totals stay valid
        r += 1
    last = r - 1
    t = r
    ws.cell(t, 1, "Total").font = Font(bold=True)
    for col in ("oh", "act", "fc", "var"):
        ws[f"{C[col]}{t}"] = f"=SUM({C[col]}{FIRST_ROW}:{C[col]}{last})"
        ws[f"{C[col]}{t}"].font = Font(bold=True)
        ws[f"{C[col]}{t}"].number_format = "#,##0"
    ws[f"{C['varp']}{t}"] = f"=IF({C['act']}{t}=0,0,{C['var']}{t}/{C['act']}{t})"
    ws[f"{C['varp']}{t}"].number_format = "0.0%"
    ws[f"{C['doc']}{t}"] = f'=IFERROR(AVERAGE({C["doc"]}{FIRST_ROW}:{C["doc"]}{last}),"")'
    ws[f"{C['doc']}{t}"].number_format = "0.0"
    for cell in ws[t]:
        cell.border = Border(top=Side(style="thin"))

    rng = lambda col: f"{C[col]}{FIRST_ROW}:{C[col]}{last}"  # noqa: E731
    ws.conditional_formatting.add(rng("risk"), CellIsRule(operator="equal", formula=['"HIGH"'], fill=PatternFill("solid", bgColor="F4B6B6")))
    ws.conditional_formatting.add(rng("risk"), CellIsRule(operator="equal", formula=['"MEDIUM"'], fill=PatternFill("solid", bgColor="FCE4A8")))
    ws.conditional_formatting.add(rng("risk"), CellIsRule(operator="equal", formula=['"LOW"'], fill=PatternFill("solid", bgColor="CDEBCB")))
    ws.conditional_formatting.add(rng("doc"), ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                              mid_type="num", mid_value=14, mid_color="FFEB84",
                                                              end_type="num", end_value=45, end_color="63BE7B"))
    ws.conditional_formatting.add(rng("varp"), FormulaRule(formula=[f"ABS({C['varp']}{FIRST_ROW})>0.25"],
                                                           font=Font(color="C00000", bold=True)))
    widths = [8, 20, 11, 7, 10, 13, 9, 9, 10, 11, 11, 13, 13, 11, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[FIRST_ROW - 1].height = 45
    ws.freeze_panes = ws.cell(FIRST_ROW, 3)
    return title, FIRST_ROW, last


def build_workbook(inv: pd.DataFrame, fva: pd.DataFrame, plants: pd.DataFrame, as_of: str) -> bytes:
    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    hold = fva[~fva["is_future"].astype(bool)].groupby(["plant_id", "sku_id"], as_index=False)[["actual", "forecast"]].sum()

    meta = []
    for p in plants.itertuples():
        title, r1, r2 = _plant_sheet(wb, p.plant_id, p.plant_name, inv, hold[hold["plant_id"] == p.plant_id])
        meta.append((p.plant_id, p.plant_name, title, r1, r2))

    sm["A1"] = "Aapurti: supply-chain summary"
    sm["A1"].font = Font(bold=True, size=15)
    sm["A2"] = f"{SYNTHETIC_NOTICE} Inventory as of {as_of}."
    sm["A2"].font = Font(italic=True, color="7F7F7F")

    labels = ["Plant", "SKUs", "On hand (units)", "Avg days of cover", "HIGH-risk SKUs",
              "Actual units (28d)", "Forecast units (28d)", "Variance (A-F)", "Variance %"]
    _header(sm, 4, labels)
    r = 5
    for pid, pname, title, r1, r2 in meta:
        q = f"'{title}'!"
        sm.cell(r, 1, f"{pid} {pname}")
        sm.cell(r, 2, f"=COUNTA({q}{C['sku']}{r1}:{C['sku']}{r2})")
        sm.cell(r, 3, f"=SUM({q}{C['oh']}{r1}:{C['oh']}{r2})")
        sm.cell(r, 4, f'=IFERROR(AVERAGE({q}{C["doc"]}{r1}:{C["doc"]}{r2}),"")')
        sm.cell(r, 5, f'=COUNTIF({q}{C["risk"]}{r1}:{C["risk"]}{r2},"HIGH")')
        sm.cell(r, 6, f"=SUM({q}{C['act']}{r1}:{C['act']}{r2})")
        sm.cell(r, 7, f"=SUM({q}{C['fc']}{r1}:{C['fc']}{r2})")
        sm.cell(r, 8, f"=F{r}-G{r}")
        sm.cell(r, 9, f"=IF(F{r}=0,0,H{r}/F{r})")
        r += 1
    first, last = 5, r - 1
    sm.cell(r, 1, "Total").font = Font(bold=True)
    for col in (2, 3, 5, 6, 7, 8):
        L = get_column_letter(col)
        sm.cell(r, col, f"=SUM({L}{first}:{L}{last})").font = Font(bold=True)
    # total average days of cover is the mean across all SKU rows, not the mean of plant means
    docs = ",".join(f"'{t}'!{C['doc']}{a}:{C['doc']}{b}" for _, _, t, a, b in meta)
    sm.cell(r, 4, f'=IFERROR(AVERAGE({docs}),"")').font = Font(bold=True)
    sm.cell(r, 9, f"=IF(F{r}=0,0,H{r}/F{r})").font = Font(bold=True)
    total_row = r
    for rr in range(first, total_row + 1):
        for col, fmt in ((2, "0"), (3, "#,##0"), (4, "0.0"), (5, "0"), (6, "#,##0"), (7, "#,##0"), (8, "#,##0"), (9, "0.0%")):
            sm.cell(rr, col).number_format = fmt

    # Pivot-style matrices (formulas over the plant sheets)
    pr = total_row + 3
    sm.cell(pr - 1, 1, "Pivot: SKU count by stockout risk x plant").font = Font(bold=True, size=12)
    _header(sm, pr, ["Risk"] + [pid for pid, *_ in meta] + ["Total"])
    for i, risk in enumerate(["HIGH", "MEDIUM", "LOW"], 1):
        rr = pr + i
        sm.cell(rr, 1, risk)
        for j, (pid, pname, title, r1, r2) in enumerate(meta, 2):
            sm.cell(rr, j, f"=COUNTIF('{title}'!{C['risk']}{r1}:{C['risk']}{r2},$A{rr})")
        sm.cell(rr, len(meta) + 2, f"=SUM(B{rr}:{get_column_letter(len(meta) + 1)}{rr})")
    risk_first, risk_last = pr + 1, pr + 3

    cats = sorted(inv["category"].unique())
    cr = pr + 7
    sm.cell(cr - 1, 1, "Pivot: on-hand units by category x plant").font = Font(bold=True, size=12)
    _header(sm, cr, ["Category"] + [pid for pid, *_ in meta] + ["Total"])
    for i, cat in enumerate(cats, 1):
        rr = cr + i
        sm.cell(rr, 1, cat)
        for j, (pid, pname, title, r1, r2) in enumerate(meta, 2):
            sm.cell(rr, j, f"=SUMIF('{title}'!{C['cat']}{r1}:{C['cat']}{r2},$A{rr},'{title}'!{C['oh']}{r1}:{C['oh']}{r2})")
            sm.cell(rr, j).number_format = "#,##0"
        sm.cell(rr, len(meta) + 2, f"=SUM(B{rr}:{get_column_letter(len(meta) + 1)}{rr})").number_format = "#,##0"

    sm.conditional_formatting.add(f"E{first}:E{last}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F4B6B6"))
    sm.conditional_formatting.add(f"I{first}:I{total_row}", FormulaRule(formula=[f"ABS(I{first})>0.25"], font=Font(color="C00000", bold=True)))
    sm.conditional_formatting.add(f"B{risk_first}:{get_column_letter(len(meta) + 1)}{risk_first}",
                                  CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", bgColor="F4B6B6")))
    for i, w in enumerate([22, 12, 16, 16, 14, 16, 18, 14, 12], 1):
        sm.column_dimensions[get_column_letter(i)].width = w
    sm.row_dimensions[4].height = 32
    sm.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
